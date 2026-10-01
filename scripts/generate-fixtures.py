#!/usr/bin/env python3
"""Synthetic fixtures. stdlib only; --producer also requires python-docx."""
import argparse
import io
import json
import struct
import zlib
from pathlib import Path
import zipfile

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Tests/SwiftDocumentsTests/Fixtures"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
def png_chunk(kind, payload):
    return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", zlib.crc32(kind + payload))

PNG = b"\x89PNG\r\n\x1a\n" + png_chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + png_chunk(b"IDAT", zlib.compress(b"\x00\x33\x88\xcc")) + png_chunk(b"IEND", b"")

def relationships(items):
    return f'<Relationships xmlns="{PKG}">' + "".join(
        f'<Relationship Id="{i}" Type="{R}/{kind}" Target="{target}"' + (' TargetMode="External"' if external else '') + '/>'
        for i, kind, target, external in items) + '</Relationships>'

def package(body, main="word/document.xml", extras=None, relations=None, macro=False):
    extras = extras or {}
    mime = "application/vnd.ms-word.document.macroEnabled.main+xml" if macro else "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
    types = f'<Types xmlns="{CT}"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/><Override PartName="/{main}" ContentType="{mime}"/></Types>'
    parts = {"[Content_Types].xml": types, "_rels/.rels": relationships([("main", "officeDocument", main, False)]),
        main: f'<w:document xmlns:w="{W}" xmlns:r="{R}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:x="urn:unknown"><w:body>{body}</w:body></w:document>', **extras}
    if relations:
        path = Path(main)
        parts[str(path.parent / "_rels" / (path.name + ".rels"))] = relationships(relations)
    return parts

def save(name, parts, compression=zipfile.ZIP_DEFLATED):
    with zipfile.ZipFile(OUT / name, "w", compression=compression) as z:
        for path, content in parts.items():
            info = zipfile.ZipInfo(path, (2026, 1, 1, 0, 0, 0))
            info.compress_type = compression
            z.writestr(info, content.encode() if isinstance(content, str) else content)

def word_root(name, body):
    return f'<w:{name} xmlns:w="{W}" xmlns:r="{R}">{body}</w:{name}>'

def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--producer", action="store_true"); args = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    simple = package('<w:p><w:r><w:t xml:space="preserve">こんにちは &amp; Swift </w:t><w:tab/><w:t>Documents</w:t><w:br/><w:t>第二行</w:t></w:r></w:p><w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:left="1440"/></w:sectPr>')
    save("minimal.docx", simple); save("stored.docx", simple, zipfile.ZIP_STORED)
    save("strict.docx", {k: v.replace(W, "http://purl.oclc.org/ooxml/wordprocessingml/main").replace(R, "http://purl.oclc.org/ooxml/officeDocument/relationships") for k, v in simple.items()})
    save("relocated.docx", package('<w:p><w:r><w:t>別の場所</w:t></w:r></w:p>', main="content/main.xml"))
    styles = word_root("styles", '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Default"/><w:sz w:val="22"/></w:rPr></w:rPrDefault></w:docDefaults><w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:rPr><w:b/></w:rPr></w:style><w:style w:type="paragraph" w:styleId="Heading"><w:name w:val="見出し"/><w:basedOn w:val="Normal"/><w:pPr><w:jc w:val="center"/><w:keepNext/></w:pPr><w:rPr><w:b/><w:i/><w:sz w:val="36"/></w:rPr></w:style><w:style w:type="character" w:styleId="Emphasis"><w:name w:val="Emphasis"/><w:rPr><w:i/></w:rPr></w:style>')
    numbering = word_root("numbering", '<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl></w:abstractNum><w:num w:numId="4"><w:abstractNumId w:val="0"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="7"/></w:lvlOverride></w:num>')
    p = '<w:p><w:pPr><w:pStyle w:val="Heading"/><w:spacing w:before="240" w:after="120"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="4"/></w:numPr></w:pPr><w:bookmarkStart w:id="0" w:name="start"/><w:r><w:t>見出し</w:t></w:r><w:r><w:rPr><w:rStyle w:val="Emphasis"/><w:b w:val="0"/><w:color w:val="FF0000"/></w:rPr><w:t>強調</w:t></w:r><w:hyperlink r:id="link" w:tooltip="案内"><w:r><w:t>リンク</w:t></w:r></w:hyperlink><w:bookmarkEnd w:id="0"/></w:p>'
    p += '<w:p><w:r><w:fldChar w:fldCharType="begin" w:fldLock="1"/></w:r><w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>3</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r><w:fldSimple w:instr="DATE"><w:r><w:t>2026/01/01</w:t></w:r></w:fldSimple></w:p>'
    p += '<w:p><w:r><w:t>変更:</w:t></w:r><w:del w:id="1" w:author="Sample" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>旧</w:delText></w:r></w:del><w:ins w:id="2" w:author="Sample"><w:r><w:t>新</w:t></w:r></w:ins><w:r><w:footnoteReference w:id="1"/><w:endnoteReference w:id="2"/></w:r><w:commentRangeStart w:id="3"/><w:r><w:t>注釈対象</w:t><w:commentReference w:id="3"/></w:r><w:commentRangeEnd w:id="3"/></w:p>'
    p += '<w:p><w:r><w:drawing><wp:inline><wp:extent cx="914400" cy="457200"/><wp:docPr id="1" name="Sample image" descr="架空の画像"/><a:graphic><a:graphicData><a:blip r:embed="image"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
    p += '<w:tbl><w:tblPr><w:tblW w:w="3000" w:type="dxa"/><w:tblBorders><w:top w:val="single" w:sz="8"/></w:tblBorders></w:tblPr><w:tblGrid><w:gridCol w:w="1500"/><w:gridCol w:w="1500"/></w:tblGrid><w:tr><w:trPr><w:tblHeader/></w:trPr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge w:val="restart"/><w:tcW w:w="3000" w:type="dxa"/><w:shd w:fill="EEEEEE"/></w:tcPr><w:p><w:r><w:t>結合</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc><w:tc><w:p><w:r><w:t>セル</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>入れ子</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:tc></w:tr></w:tbl>'
    p += '<w:sectPr><w:headerReference w:type="default" r:id="header"/><w:footerReference w:type="default" r:id="footer"/><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:left="1800"/><w:cols w:num="2"/><w:titlePg/></w:sectPr>'
    extras = {"word/styles.xml": styles, "word/numbering.xml": numbering,
        "word/header.xml": word_root("hdr", '<w:p><w:r><w:t>ヘッダー</w:t></w:r></w:p>'),
        "word/footer.xml": word_root("ftr", '<w:p><w:r><w:t>フッター</w:t></w:r></w:p>'),
        "word/footnotes.xml": word_root("footnotes", '<w:footnote w:id="-1" w:type="separator"><w:p><w:r><w:separator/></w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>脚注</w:t></w:r></w:p></w:footnote>'),
        "word/endnotes.xml": word_root("endnotes", '<w:endnote w:id="2"><w:p><w:r><w:t>文末脚注</w:t></w:r></w:p></w:endnote>'),
        "word/comments.xml": word_root("comments", '<w:comment w:id="3" w:author="Sample"><w:p><w:r><w:t>コメント</w:t></w:r></w:p></w:comment>'),
        "word/media/sample.png": PNG,
        "docProps/core.xml": '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/"><dc:title>架空の文書</dc:title><dc:creator>Sample Author</dc:creator><dcterms:created>2026-01-01T00:00:00Z</dcterms:created></cp:coreProperties>',
        "docProps/app.xml": '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Application>Fixture Generator</Application><Pages>2</Pages><Words>20</Words></Properties>',
        "docProps/custom.xml": '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property name="Project"><vt:lpwstr>Sample</vt:lpwstr></property></Properties>'}
    relations = [(name, kind, target, False) for name, kind, target in [
        ("styles", "styles", "styles.xml"), ("numbering", "numbering", "numbering.xml"), ("header", "header", "header.xml"), ("footer", "footer", "footer.xml"),
        ("footnotes", "footnotes", "footnotes.xml"), ("endnotes", "endnotes", "endnotes.xml"), ("comments", "comments", "comments.xml"), ("image", "image", "media/sample.png")]] + [("link", "hyperlink", "https://example.com/guide", True)]
    parts = package(p, extras=extras, relations=relations)
    parts["_rels/.rels"] = relationships([("main", "officeDocument", "word/document.xml", False),
        ("core", "metadata/core-properties", "docProps/core.xml", False), ("app", "extended-properties", "docProps/app.xml", False), ("custom", "custom-properties", "docProps/custom.xml", False)])
    save("features.docx", parts)
    save("macro.docm", package('<w:p><w:r><w:t>マクロは実行しない</w:t></w:r></w:p>', extras={"word/vbaProject.bin": b"SYNTHETIC-NOT-EXECUTABLE"}, relations=[("vba", "vbaProject", "vbaProject.bin", False)], macro=True))
    save("writer.odt", {"mimetype": "application/vnd.oasis.opendocument.text", "content.xml": "<document/>"})
    def iwa_type(type_id):
        def varint(n):
            out = bytearray()
            while n > 127: out.append((n & 127) | 128); n >>= 7
            out.append(n); return bytes(out)
        info = b"\x08" + varint(type_id) + b"\x18\x00"
        header = b"\x08\x01\x12" + varint(len(info)) + info
        payload = varint(len(header)) + header
        compressed = varint(len(payload)) + bytes([(len(payload) - 1) << 2]) + payload
        return b"\x00" + len(compressed).to_bytes(3, "little") + compressed
    save("design.pages", {"Index/Document.iwa": iwa_type(10000)})
    save("spreadsheet.iwork", {"Index/Document.iwa": iwa_type(1)})
    save("unknown.docx", package('<w:sdt><w:sdtContent><w:p><w:r><w:t>表示内容</w:t></w:r></w:p></w:sdtContent></w:sdt><w:p><w:r><w:rPr><w:emboss/></w:rPr><w:t>既知</w:t><x:widget><x:t>未知</x:t></x:widget></w:r><m:oMath><m:r><m:t>x+y</m:t></m:r></m:oMath></w:p><w:altChunk r:id="html"/>', extras={"word/chunk.html": "<p>html</p>"}, relations=[("html", "aFChunk", "chunk.html", False)]))
    if args.producer:
        from docx import Document
        from docx.shared import Inches, Pt
        d = Document(); d.core_properties.author = "Sample Author"; d.core_properties.last_modified_by = "Sample Author"
        d.core_properties.title = "Producer fixture"
        d.add_heading("SwiftDocuments 検証", 1)
        p = d.add_paragraph("日本語と emoji 📄 ")
        r = p.add_run("太字"); r.bold = True; r.font.size = Pt(14)
        p.add_run(" / "); p.add_run("斜体").italic = True
        d.add_paragraph("項目1", style="List Number"); d.add_paragraph("項目2", style="List Number")
        table = d.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [["項目", "値"], ["売上", "1,234"]]):
            for cell, value in zip(row.cells, values): cell.text = value
        d.add_picture(io.BytesIO(PNG), width=Inches(0.25))
        d.sections[0].header.paragraphs[0].text = "Producer header"; d.sections[0].footer.paragraphs[0].text = "Producer footer"
        d.save(OUT / "python-docx.docx")
        expected = {"paragraphs": [p.text for p in d.paragraphs], "tables": [[[c.text for c in r.cells] for r in t.rows] for t in d.tables], "header": "Producer header", "footer": "Producer footer"}
        (OUT / "producer-oracle.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2) + "\n")

if __name__ == "__main__": main()
